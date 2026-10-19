#!/usr/bin/env python3
"""Builds template/sample-data.csv and template/nyyo-sheets-kit-template.xlsx from the sample rows in
apps-script/Sheets.gs, so the template, the kit's "Load sample data" and the tests always hold the same rows.
   python3 scripts/make_template.py      (needs openpyxl; run by scripts/build_dist.sh)"""
import csv, json, re, sys, zipfile, io
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
src = (ROOT / "apps-script/Sheets.gs").read_text(encoding="utf-8")
block = re.search(r"var NYYO_SAMPLE_ = (\[.*?\n\]);", src, re.S).group(1)
rows = json.loads(block.replace("'", '"'))
cfg = (ROOT / "apps-script/Config.gs").read_text(encoding="utf-8")
headers = json.loads(re.search(r"headers: (\[[^\]]*\])", cfg).group(1).replace("'", '"'))
data = [r + [""] * (len(headers) - len(r)) for r in rows]

out = ROOT / "template"; out.mkdir(exist_ok=True)
with open(out / "sample-data.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh, lineterminator="\n"); w.writerow(headers); w.writerows(data)

from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook(); ws = wb.active; ws.title = "Data"
ws.append(headers); [ws.append(r) for r in data]
for c in ws[1]: c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.number_format = "@"   # keep the messy text exactly as typed
ws.freeze_panes = "A2"
for i, h in enumerate(headers): ws.column_dimensions[chr(65 + i)].width = max(12, len(h) + 4)
hw = wb.create_sheet("How to use")
for line in [
    "Nyyo Sheets kit - sample template",
    "",
    "1. In Google Sheets: File > Import > Upload this file > Replace spreadsheet (or Insert new sheet(s)).",
    "2. Install the kit in that spreadsheet (INSTALL.md), then menu Nyyo kit > 1. Set up.",
    "3. Set up keeps the rows already in the Data tab and adds the Clean, Dashboard and Change log tabs.",
    "   Its 'How to use' tab then replaces this one with the full guide.",
    "",
    "Questions? Ask in your order chat.",
]:
    hw.append([line])
hw["A1"].font = Font(bold=True); hw.column_dimensions["A"].width = 100
wb.properties.creator = "Nyyo Automation"; wb.properties.lastModifiedBy = "Nyyo Automation"
import datetime
fixed = datetime.datetime(2026, 10, 19, 0, 0, 0)
wb.properties.created = fixed; wb.properties.modified = fixed
buf = io.BytesIO(); wb.save(buf)
# re-zip with fixed timestamps so the same input gives the same file
src_zip = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
with zipfile.ZipFile(out / "nyyo-sheets-kit-template.xlsx", "w", zipfile.ZIP_DEFLATED) as z:
    for name in sorted(src_zip.namelist()):
        data_bytes = src_zip.read(name)
        if name == "docProps/core.xml":
            data_bytes = re.sub(rb"(<dcterms:(?:modified|created)[^>]*>)[^<]*(</dcterms:(?:modified|created)>)",
                                rb"\g<1>2026-10-19T00:00:00Z\g<2>", data_bytes)
        zi = zipfile.ZipInfo(name, date_time=(2026, 10, 19, 0, 0, 0)); zi.compress_type = zipfile.ZIP_DEFLATED
        z.writestr(zi, data_bytes)
print(f"template: {len(data)} sample rows, {len(headers)} columns -> template/sample-data.csv, template/nyyo-sheets-kit-template.xlsx")
